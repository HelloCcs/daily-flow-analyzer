from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analysis import CATEGORIES, PROGRESSION_PAIRS, WINDOWS, build_analysis_rows
from .history_store import HistoryStore
from .providers import LimitRecord, MarketRecord


BASE_COLUMNS = [
    "序号", "代码", "名称", "交易日期", "涨跌幅(%)", "收盘价(元)", "总市值(亿元)", "成交额(万元)",
    "主力资金流入比率(%)", "主力净额占总市值(%)", "是否涨停", "涨停幅度(%)", "几天几板",
    "连续涨停板数", "封单资金(万元)", "首次封板时间", "最后封板时间", "炸板次数",
]
DAILY_COLUMNS = [
    value
    for _, label in CATEGORIES
    for value in (
        f"今日{label}净额(万元)", f"昨日{label}净额(万元)", f"{label}净额增减(万元)",
        f"{label}净额增长率(%)", f"{label}变化状态",
    )
]
PERIOD_COLUMNS = []
for _window in WINDOWS:
    for _, _label in CATEGORIES:
        PERIOD_COLUMNS.extend((f"{_window}日{_label}净额合计(万元)", f"{_window}日{_label}较上期差值(万元)"))
    PERIOD_COLUMNS.extend((
        f"{_window}日含今日主力资金流入比率(%)",
        f"{_window}日不含今日主力资金流入比率(%)",
        f"{_window}日比率分差(百分点)",
    ))
ALL_COLUMNS = BASE_COLUMNS + DAILY_COLUMNS + PERIOD_COLUMNS

SUMMARY_COLUMNS = [
    "序号", "代码", "名称", "交易日期", "涨跌幅(%)", "收盘价(元)", "总市值(亿元)",
    "今日主力资金流入比率(%)", "昨日主力资金流入比率(%)", "日比率分差(百分点)",
]
for _window in WINDOWS:
    SUMMARY_COLUMNS.extend((
        f"{_window}日含今日主力资金流入比率(%)",
        f"{_window}日不含今日主力资金流入比率(%)",
        f"{_window}日比率分差(百分点)",
    ))
MARKET_RATIO_COLUMNS = [
    "今日主力净额占总市值(%)", "今日成交额占总市值(%)",
]
for _window in WINDOWS:
    MARKET_RATIO_COLUMNS.extend((
        f"{_window}日含今日主力净额占总市值(%)",
        f"{_window}日含今日成交额占总市值(%)",
    ))
SUMMARY_COLUMNS.extend(MARKET_RATIO_COLUMNS)
SUMMARY_COLUMNS.extend((
    "是否涨停", "涨停幅度(%)", "几天几板", "连续涨停板数", "封单资金(万元)",
    "首次封板时间", "最后封板时间", "炸板次数",
))

SORT_OPTIONS = {
    "日比率分差(百分点)": "日比率分差(百分点)",
    "今日主力资金流入比率(%)": "今日主力资金流入比率(%)",
    "3日主力资金流入比率(%)": "3日含今日主力资金流入比率(%)",
    "5日主力资金流入比率(%)": "5日含今日主力资金流入比率(%)",
    "7日主力资金流入比率(%)": "7日含今日主力资金流入比率(%)",
    "10日主力资金流入比率(%)": "10日含今日主力资金流入比率(%)",
    "逐步增减值(百分点)": "逐步增减值(百分点)",
    "今日主力净额(万元)": "今日主力净额(万元)",
    "3日主力净额(万元)": "3日含今日主力净额(万元)",
    "5日主力净额(万元)": "5日含今日主力净额(万元)",
    "7日主力净额(万元)": "7日含今日主力净额(万元)",
    "10日主力净额(万元)": "10日含今日主力净额(万元)",
    "涨跌幅(%)": "涨跌幅(%)",
    "总市值(亿元)": "总市值(亿元)",
    "今日主力净额占总市值(%)": "今日主力净额占总市值(%)",
    "3日主力净额占总市值(%)": "3日含今日主力净额占总市值(%)",
    "5日主力净额占总市值(%)": "5日含今日主力净额占总市值(%)",
    "7日主力净额占总市值(%)": "7日含今日主力净额占总市值(%)",
    "10日主力净额占总市值(%)": "10日含今日主力净额占总市值(%)",
    "当日倍量(倍)": "倍量1(倍)",
    "代码": "代码",
}
DEFAULT_SORT_OPTION = "日比率分差(百分点)"


def _sort_analysis_rows(rows: list[dict], sort_option: str, descending: bool) -> list[dict]:
    key = SORT_OPTIONS.get(sort_option, SORT_OPTIONS[DEFAULT_SORT_OPTION])
    present = [row for row in rows if row.get(key) is not None]
    missing = [row for row in rows if row.get(key) is None]
    present.sort(key=lambda row: row.get(key), reverse=descending)
    return present + missing


def _component_columns(window: int | None) -> list[tuple[str, str]]:
    columns = [
        ("序号", "序号"), ("代码", "代码"), ("名称", "名称"), ("交易日期", "交易日期"),
        ("收盘价(元)", "收盘价(元)"), ("总市值(亿元)", "总市值(亿元)"),
    ]
    for _, label in CATEGORIES:
        if window is None:
            columns.extend((
                ("今日净额(万元)", f"今日{label}净额(万元)"),
                ("昨日净额(万元)", f"昨日{label}净额(万元)"),
                ("增减量(万元)", f"{label}净额增减(万元)"),
                ("今日资金流入比率(%)", f"今日{label}资金流入比率(%)"),
                ("昨日资金流入比率(%)", f"昨日{label}资金流入比率(%)"),
                ("比率分差(百分点)", f"{label}资金流入比率分差(百分点)"),
            ))
        else:
            columns.extend((
                ("含今日净额(万元)", f"{window}日含今日{label}净额(万元)"),
                ("不含今日净额(万元)", f"{window}日不含今日{label}净额(万元)"),
                ("增减量(万元)", f"{window}日{label}增减量(万元)"),
                ("含今日资金流入比率(%)", f"{window}日含今日{label}资金流入比率(%)"),
                ("不含今日资金流入比率(%)", f"{window}日不含今日{label}资金流入比率(%)"),
                ("比率分差(百分点)", f"{window}日{label}资金流入比率分差(百分点)"),
            ))
    return columns


def _progression_columns(scope: str) -> list[tuple[str, str]]:
    columns = [
        ("序号", "序号"), ("代码", "代码"), ("名称", "名称"), ("交易日期", "交易日期"),
        ("收盘价(元)", "收盘价(元)"), ("总市值(亿元)", "总市值(亿元)"),
    ]
    for _, label in CATEGORIES:
        for window in WINDOWS:
            columns.extend((
                (f"{window}日净额(万元)", f"{window}日{scope}{label}净额(万元)"),
                (f"{window}日资金流入比率(%)", f"{window}日{scope}{label}资金流入比率(%)"),
            ))
        for shorter, longer in PROGRESSION_PAIRS:
            columns.extend((
                (f"{shorter}→{longer}日增加量(万元)", f"{scope}{shorter}至{longer}日{label}增加量(万元)"),
                (f"{shorter}→{longer}日比率变化(百分点)", f"{scope}{shorter}至{longer}日{label}比率变化(百分点)"),
            ))
    return columns


def _observation_columns() -> list[tuple[str, str]]:
    columns = [
        ("序号", "序号"), ("代码", "代码"), ("名称", "名称"), ("交易日期", "交易日期"),
        ("收盘价(元)", "收盘价(元)"), ("总市值(亿元)", "总市值(亿元)"),
    ]
    for window in WINDOWS:
        columns.extend((
            (f"{window}日主力净额(万元)", f"{window}日含今日主力净额(万元)"),
            (f"{window}日成交额(万元)", f"{window}日含今日成交额(万元)"),
            (f"{window}日主力资金流入比率(%)", f"{window}日含今日主力资金流入比率(%)"),
        ))
    # 直接复用前面各区间已经计算好的比率，集中横向展示，方便比较。
    for window in WINDOWS:
        columns.append((
            f"{window}日主力资金流入比率(%)",
            f"{window}日含今日主力资金流入比率(%)",
        ))
    for longer, shorter in PROGRESSION_PAIRS:
        columns.append((
            f"{longer}→{shorter}日比率变化(百分点)",
            f"{longer}至{shorter}日比率变化(百分点)",
        ))
    columns.extend((
        ("逐步增减值(百分点)", "逐步增减值(百分点)"),
        ("是否逐步增长", "是否逐步增长"),
        ("变化观察", "周期资金变化观察"),
    ))
    columns.extend((header, header) for header in MARKET_RATIO_COLUMNS)
    return columns


def _volume_columns(dates: list[date]) -> list[tuple[str, str]]:
    columns = [
        ("序号", "序号"), ("代码", "代码"), ("名称", "名称"), ("交易日期", "交易日期"),
        ("收盘价(元)", "收盘价(元)"), ("总市值(亿元)", "总市值(亿元)"),
    ]
    for index, day in enumerate(dates[:10], start=1):
        label = day.strftime("%m-%d")
        columns.extend((
            (f"{label}成交量(万手)", f"成交量{index}(万手)"),
            (f"{label}倍量(倍)", f"倍量{index}(倍)"),
        ))
    return columns

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
RED = "C00000"
GREEN = "008000"
PALE_RED = "FCE4D6"
PALE_GREEN = "E2F0D9"
GRAY = "E7E6E6"
THIN = Side(style="thin", color="B7C9D6")


def _group_ranges() -> list[tuple[str, int, int]]:
    groups = [("基础信息 / 市值 / 涨停", 1, len(BASE_COLUMNS))]
    start = len(BASE_COLUMNS) + 1
    for _, label in CATEGORIES:
        groups.append((f"最近交易日 vs 上一交易日 - {label}", start, start + 4))
        start += 5
    for window in WINDOWS:
        groups.append((f"{window}日滚动合计、主力流入比率及较上一滚动窗口差值", start, start + 10))
        start += 11
    return groups


def _style_analysis_sheet(sheet, rows: list[dict]) -> None:
    for title, start, end in _group_ranges():
        if end > start:
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        cell = sheet.cell(1, start, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(start, end + 1):
            sheet.cell(1, column).fill = PatternFill("solid", fgColor=NAVY)
            sheet.cell(1, column).border = Border(bottom=THIN)
    for column, header in enumerate(ALL_COLUMNS, start=1):
        cell = sheet.cell(2, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 42
    # Keep the complete stock identity block visible while horizontally reviewing or sorting metrics.
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(ALL_COLUMNS))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"

    widths = {"A": 7, "B": 13, "C": 14, "D": 12}
    for column in range(1, len(ALL_COLUMNS) + 1):
        letter = get_column_letter(column)
        header = ALL_COLUMNS[column - 1]
        width = widths.get(letter, 14)
        if "状态" in header or "时间" in header or "几天几板" in header:
            width = 13
        elif "增长率" in header or "比率" in header or "涨跌幅" in header:
            width = 15
        elif "净额" in header or "差值" in header or "市值" in header or "成交额" in header:
            width = 17
        sheet.column_dimensions[letter].width = width

    last_row = len(rows) + 2
    percent_headers = {"涨跌幅(%)", "主力资金流入比率(%)", "主力净额占总市值(%)", "涨停幅度(%)"}
    percent_headers.update(f"{label}净额增长率(%)" for _, label in CATEGORIES)
    percent_headers.update(f"{window}日主力资金流入比率(%)" for window in WINDOWS)
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, header in enumerate(ALL_COLUMNS, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if header in {"代码", "名称", "几天几板"} else "right")
            if header == "交易日期":
                cell.number_format = "yyyy-mm-dd"
            elif header in percent_headers:
                cell.number_format = "0.00"
            elif isinstance(cell.value, (int, float)) and header != "序号":
                cell.number_format = "#,##0.00;[Green]-#,##0.00;-"
            if header == "是否涨停" and cell.value == "是":
                cell.fill = PatternFill("solid", fgColor=PALE_RED)
                cell.font = Font(color=RED, bold=True)

    color_headers = [
        header for header in ALL_COLUMNS
        if "增减(万元)" in header or "增长率(%)" in header or "较上期差值(万元)" in header or "百分点" in header or header == "涨跌幅(%)"
    ]
    for header in color_headers:
        column = ALL_COLUMNS.index(header) + 1
        letter = get_column_letter(column)
        for row_index in range(3, last_row + 1):
            sheet.cell(row_index, column).number_format = "[Red]#,##0.00;[Green]-#,##0.00;-"
        cell_range = f"{letter}3:{letter}{last_row}"
        thresholds = (20, 50, 100) if "%" in header or "百分点" in header else (1000, 10000, 50000)
        positive = ((thresholds[2], "9C0006"), (thresholds[1], "E06666"), (thresholds[0], "EA9999"), (0, "F4CCCC"))
        negative = ((-thresholds[2], "274E13"), (-thresholds[1], "6AA84F"), (-thresholds[0], "93C47D"), (0, "D9EAD3"))
        for value, color in positive:
            operator = "greaterThan" if value == 0 else "greaterThanOrEqual"
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )
        for value, color in negative:
            operator = "lessThan" if value == 0 else "lessThanOrEqual"
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )


def _style_summary_sheet(sheet, rows: list[dict]) -> None:
    groups = [("基础信息", 1, 7), ("最近交易日 vs 上一交易日", 8, 10)]
    start = 11
    for window in WINDOWS:
        groups.append((f"{window}日：含今日 vs 不含今日", start, start + 2))
        start += 3
    groups.append(("净额及成交额占报告日总市值", start, start + len(MARKET_RATIO_COLUMNS) - 1))
    start += len(MARKET_RATIO_COLUMNS)
    groups.append(("涨停情况", start, start + 7))
    for title, start, end in groups:
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        for column in range(start, end + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        sheet.cell(1, start, title)
    for column, header in enumerate(SUMMARY_COLUMNS, start=1):
        cell = sheet.cell(2, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        width = 14
        if column == 1:
            width = 7
        elif column == 2:
            width = 13
        elif column == 3:
            width = 14
        elif column == 4:
            width = 12
        elif "资金流入比率" in header or "分差" in header:
            width = 18
        elif "时间" in header:
            width = 14
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 48
    # Keep the complete stock identity block visible while reviewing any sorted metric.
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(SUMMARY_COLUMNS))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    last_row = len(rows) + 2
    difference_headers = [header for header in SUMMARY_COLUMNS if "分差" in header]
    ratio_headers = [header for header in SUMMARY_COLUMNS if "比率" in header or "分差" in header or header in {"涨跌幅(%)", "涨停幅度(%)"}]
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, header in enumerate(SUMMARY_COLUMNS, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if header in {"代码", "名称", "几天几板"} else "right")
            if header == "交易日期":
                cell.number_format = "yyyy-mm-dd"
            elif header in ratio_headers:
                cell.number_format = "0.00"
            elif isinstance(cell.value, (int, float)) and header != "序号":
                cell.number_format = "#,##0.00;[Green]-#,##0.00;-"
            if header == "是否涨停" and cell.value == "是":
                cell.fill = PatternFill("solid", fgColor=PALE_RED)
                cell.font = Font(color=RED, bold=True)
    for header in difference_headers:
        column = SUMMARY_COLUMNS.index(header) + 1
        letter = get_column_letter(column)
        for row_index in range(3, last_row + 1):
            sheet.cell(row_index, column).number_format = "[Red]0.00;[Green]-0.00;-"
        cell_range = f"{letter}3:{letter}{last_row}"
        for value, color, operator in (
            (20, "9C0006", "greaterThanOrEqual"),
            (5, "EA9999", "greaterThanOrEqual"),
            (0, "F4CCCC", "greaterThan"),
            (-20, "274E13", "lessThanOrEqual"),
            (-5, "93C47D", "lessThanOrEqual"),
            (0, "D9EAD3", "lessThan"),
        ):
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )


def _style_component_sheet(sheet, rows: list[dict], window: int | None) -> None:
    columns = _component_columns(window)
    period_title = "最近交易日 vs 上一交易日" if window is None else f"{window}日：含今日 vs 不含今日"
    groups = [("基础信息", 1, 6)]
    start = 7
    for _, label in CATEGORIES:
        groups.append((f"{period_title} - {label}", start, start + 5))
        start += 6
    for title, group_start, group_end in groups:
        sheet.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=group_end)
        for column in range(group_start, group_end + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        sheet.cell(1, group_start, title)
    for column, (display, _) in enumerate(columns, start=1):
        cell = sheet.cell(2, column, display)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        sheet.column_dimensions[get_column_letter(column)].width = 7 if column == 1 else 13 if column == 2 else 14 if column == 3 else 12 if column == 4 else 13 if column == 5 else 16 if column == 6 else 17
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 42
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    last_row = len(rows) + 2
    change_columns = []
    for column, (display, _) in enumerate(columns, start=1):
        if display in {"增减量(万元)", "比率分差(百分点)"}:
            change_columns.append((column, display))
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, (display, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if column in {2, 3} else "right")
            if column == 4:
                cell.number_format = "yyyy-mm-dd"
            elif "资金流入比率" in display or display == "比率分差(百分点)":
                cell.number_format = "[Red]0.00;[Green]-0.00;-"
            elif isinstance(cell.value, (int, float)) and column != 1:
                cell.number_format = "#,##0.00;[Green]-#,##0.00;-"
    for column, display in change_columns:
        letter = get_column_letter(column)
        cell_range = f"{letter}3:{letter}{last_row}"
        thresholds = (1, 3, 8) if display == "比率分差(百分点)" else (1000, 10000, 50000)
        rules = (
            (thresholds[2], "9C0006", "greaterThanOrEqual"),
            (thresholds[1], "E06666", "greaterThanOrEqual"),
            (thresholds[0], "EA9999", "greaterThanOrEqual"),
            (0, "F4CCCC", "greaterThan"),
            (-thresholds[2], "274E13", "lessThanOrEqual"),
            (-thresholds[1], "6AA84F", "lessThanOrEqual"),
            (-thresholds[0], "93C47D", "lessThanOrEqual"),
            (0, "D9EAD3", "lessThan"),
        )
        for value, color, operator in rules:
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )


def _style_progression_sheet(sheet, rows: list[dict], scope: str) -> None:
    columns = _progression_columns(scope)
    groups = [(f"基础信息 / {scope}", 1, 6)]
    start = 7
    group_width = len(WINDOWS) * 2 + len(PROGRESSION_PAIRS) * 2
    for _, label in CATEGORIES:
        groups.append((f"{scope}：10日 → 7日 → 5日 → 3日 - {label}", start, start + group_width - 1))
        start += group_width
    for title, group_start, group_end in groups:
        sheet.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=group_end)
        for column in range(group_start, group_end + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        sheet.cell(1, group_start, title)
    for column, (display, _) in enumerate(columns, start=1):
        cell = sheet.cell(2, column, display)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        sheet.column_dimensions[get_column_letter(column)].width = 7 if column == 1 else 13 if column == 2 else 14 if column == 3 else 12 if column == 4 else 13 if column == 5 else 16 if column == 6 else 18
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 48
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    last_row = len(rows) + 2
    change_columns = []
    for column, (display, _) in enumerate(columns, start=1):
        if "增加量" in display or "比率变化" in display:
            change_columns.append((column, display))
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, (display, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if column in {2, 3} else "right")
            if column == 4:
                cell.number_format = "yyyy-mm-dd"
            elif "资金流入比率" in display or "比率变化" in display:
                cell.number_format = "[Red]0.00;[Green]-0.00;-"
            elif isinstance(cell.value, (int, float)) and column != 1:
                cell.number_format = "#,##0.00;[Green]-#,##0.00;-"
    for column, display in change_columns:
        letter = get_column_letter(column)
        cell_range = f"{letter}3:{letter}{last_row}"
        thresholds = (1, 3, 8) if "比率变化" in display else (1000, 10000, 50000)
        rules = (
            (thresholds[2], "9C0006", "greaterThanOrEqual"),
            (thresholds[1], "E06666", "greaterThanOrEqual"),
            (thresholds[0], "EA9999", "greaterThanOrEqual"),
            (0, "F4CCCC", "greaterThan"),
            (-thresholds[2], "274E13", "lessThanOrEqual"),
            (-thresholds[1], "6AA84F", "lessThanOrEqual"),
            (-thresholds[0], "93C47D", "lessThanOrEqual"),
            (0, "D9EAD3", "lessThan"),
        )
        for value, color, operator in rules:
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )


def _style_observation_sheet(sheet, rows: list[dict]) -> None:
    columns = _observation_columns()
    groups = [("基础信息", 1, 6)]
    start = 7
    for window in WINDOWS:
        groups.append((f"{window}日区间资金观察", start, start + 2))
        start += 3
    groups.append(("10/7/5/3日比率、比率变化与逐步增长观察", start, start + 9))
    start += 10
    groups.append(("净额及成交额占报告日总市值", start, start + len(MARKET_RATIO_COLUMNS) - 1))
    for title, group_start, group_end in groups:
        sheet.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=group_end)
        for column in range(group_start, group_end + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        sheet.cell(1, group_start, title)
    for column, (display, _) in enumerate(columns, start=1):
        cell = sheet.cell(2, column, display)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        sheet.column_dimensions[get_column_letter(column)].width = 7 if column == 1 else 13 if column == 2 else 14 if column == 3 else 12 if column == 4 else 13 if column == 5 else 16 if column == 6 else 18
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 50
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    last_row = len(rows) + 2
    ratio_columns = []
    for column, (display, _) in enumerate(columns, start=1):
        if "比率" in display or display == "逐步增减值(百分点)":
            ratio_columns.append(column)
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, (display, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if column in {2, 3} or display == "变化观察" else "right")
            if column == 4:
                cell.number_format = "yyyy-mm-dd"
            elif "比率" in display or display == "逐步增减值(百分点)":
                cell.number_format = "[Red]0.00;[Green]-0.00;-"
            elif isinstance(cell.value, (int, float)) and column != 1:
                cell.number_format = "#,##0.00;[Green]-#,##0.00;-"
            if display == "是否逐步增长":
                if cell.value == "是":
                    cell.fill = PatternFill("solid", fgColor=PALE_RED)
                    cell.font = Font(color=RED, bold=True)
                elif cell.value == "否":
                    cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
                    cell.font = Font(color=GREEN, bold=True)
    for column in ratio_columns:
        letter = get_column_letter(column)
        cell_range = f"{letter}3:{letter}{last_row}"
        for value, color, operator in (
            (3, "9C0006", "greaterThanOrEqual"),
            (1, "E06666", "greaterThanOrEqual"),
            (0, "F4CCCC", "greaterThan"),
            (-3, "274E13", "lessThanOrEqual"),
            (-1, "6AA84F", "lessThanOrEqual"),
            (0, "D9EAD3", "lessThan"),
        ):
            sheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator=operator, formula=[str(value)], fill=PatternFill("solid", fgColor=color), stopIfTrue=True),
            )


def _style_volume_sheet(sheet, rows: list[dict], dates: list[date]) -> None:
    columns = _volume_columns(dates)
    groups = [("基础信息", 1, 6)]
    start = 7
    for day in dates[:10]:
        groups.append((f"{day:%Y-%m-%d} 成交量与前一交易日比较", start, start + 1))
        start += 2
    for title, group_start, group_end in groups:
        if group_end > group_start:
            sheet.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=group_end)
        for column in range(group_start, group_end + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN)
        sheet.cell(1, group_start, title)
    for column, (display, _) in enumerate(columns, start=1):
        cell = sheet.cell(2, column, display)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        sheet.column_dimensions[get_column_letter(column)].width = 7 if column == 1 else 13 if column == 2 else 14 if column == 3 else 12 if column == 4 else 13 if column == 5 else 16 if column == 6 else 17
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 42
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(rows) + 2}"
    sheet.sheet_view.showGridLines = False
    last_row = len(rows) + 2
    for row_index in range(3, last_row + 1):
        if row_index % 2:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
        for column, (display, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center", horizontal="left" if column in {2, 3} else "right")
            if column == 4:
                cell.number_format = "yyyy-mm-dd"
            elif "倍量" in display:
                cell.number_format = "0.00"
                if isinstance(cell.value, (int, float)):
                    if cell.value >= 3:
                        cell.fill = PatternFill("solid", fgColor="000000")
                        cell.font = Font(color=WHITE, bold=True)
                    elif cell.value >= 2:
                        cell.fill = PatternFill("solid", fgColor="7030A0")
                        cell.font = Font(color=WHITE, bold=True)
                    elif cell.value >= 1.5:
                        cell.fill = PatternFill("solid", fgColor="5B9BD5")
                        cell.font = Font(color=WHITE, bold=True)
            elif isinstance(cell.value, (int, float)) and column != 1:
                cell.number_format = "#,##0.00"


def _write_simple_sheet(sheet, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, _excel_safe(value))
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    sheet.sheet_view.showGridLines = False
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1] if widths else 18


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_report(
    store: HistoryStore,
    report_date: date,
    market: dict[str, MarketRecord],
    limits: dict[str, LimitRecord],
    output_path: str | Path,
    progress: Callable[[str], None] = lambda _: None,
    sort_option: str = DEFAULT_SORT_OPTION,
    sort_descending: bool = True,
    volume_history: dict[str, dict[str, float]] | None = None,
) -> Path:
    dates, flow_rows = store.load_window(report_date, 11)
    if not dates or dates[0] != report_date:
        raise ValueError(f"历史库中没有 {report_date} 的资金表。")
    progress("正在计算昨日对比和 3/5/7/10 日滚动指标...")
    rows = build_analysis_rows(dates, flow_rows, market, limits, volume_history)
    if len(rows) < 5000:
        raise ValueError(f"报告日仅有 {len(rows)} 只股票，拒绝生成不完整报表。")
    rows = _sort_analysis_rows(rows, sort_option, sort_descending)
    sort_direction = "降序" if sort_descending else "升序"
    progress(f"统一排序：{sort_option}（{sort_direction}），所有分析页采用同一股票顺序。")
    code_order = {row["代码"]: index for index, row in enumerate(rows)}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "资金比率简表"
    for row_index, item in enumerate(rows, start=3):
        for column, header in enumerate(SUMMARY_COLUMNS, start=1):
            summary.cell(row_index, column, _excel_safe(item.get(header)))
    _style_summary_sheet(summary, rows)

    observation_sheet = workbook.create_sheet("周期资金观察", 0)
    observation_columns = _observation_columns()
    for row_index, item in enumerate(rows, start=3):
        for column, (_, key) in enumerate(observation_columns, start=1):
            observation_sheet.cell(row_index, column, _excel_safe(item.get(key)))
    _style_observation_sheet(observation_sheet, rows)

    volume_sheet = workbook.create_sheet("10日成交量倍量", 1)
    volume_columns = _volume_columns(dates)
    for row_index, item in enumerate(rows, start=3):
        for column, (_, key) in enumerate(volume_columns, start=1):
            volume_sheet.cell(row_index, column, _excel_safe(item.get(key)))
    _style_volume_sheet(volume_sheet, rows, dates)

    for window, sheet_name in ((None, "日资金成分对比"), (10, "10日成分对比"), (7, "7日成分对比"), (5, "5日成分对比"), (3, "3日成分对比")):
        component_sheet = workbook.create_sheet(sheet_name)
        columns = _component_columns(window)
        for row_index, item in enumerate(rows, start=3):
            for column, (_, key) in enumerate(columns, start=1):
                component_sheet.cell(row_index, column, _excel_safe(item.get(key)))
        _style_component_sheet(component_sheet, rows, window)

    for scope, sheet_name in (("含今日", "含今日周期递进"), ("不含今日", "不含今日周期递进")):
        progression_sheet = workbook.create_sheet(sheet_name)
        columns = _progression_columns(scope)
        for row_index, item in enumerate(rows, start=3):
            for column, (_, key) in enumerate(columns, start=1):
                progression_sheet.cell(row_index, column, _excel_safe(item.get(key)))
        _style_progression_sheet(progression_sheet, rows, scope)

    latest_rows = [row for row in flow_rows if row["trade_date"] == report_date.isoformat()]
    detail = workbook.create_sheet("当日源数据")
    detail_headers = ["交易日期", "代码", "名称", "涨跌幅(%)", "成交额(万元)", "主力净额(万元)",
                      "主力流入比率(%)", "特大单净额(万元)", "大单净额(万元)", "中单净额(万元)", "小单净额(万元)"]
    latest_rows.sort(key=lambda row: code_order.get(row["code"], len(code_order)))
    detail_rows = [[row[key] for key in ("trade_date", "code", "name", "pct_change", "turnover", "main_net",
                                          "main_ratio", "super_net", "big_net", "mid_net", "small_net")]
                   for row in latest_rows]
    _write_simple_sheet(detail, detail_headers, detail_rows, [12, 13, 14] + [17] * 8)

    limit_sheet = workbook.create_sheet("涨停明细")
    limit_headers = ["代码", "名称", "涨停幅度(%)", "几天几板", "连续涨停板数", "封单资金(万元)",
                     "首次封板时间", "最后封板时间", "炸板次数"]
    sorted_limits = sorted(limits.values(), key=lambda item: code_order.get(item.code, len(code_order)))
    limit_rows = [[item.code, item.name, item.limit_pct, item.board_text, item.consecutive_boards,
                   item.seal_amount_yuan / 10000 if item.seal_amount_yuan is not None else None,
                   item.first_seal_time, item.last_seal_time, item.break_count]
                  for item in sorted_limits]
    _write_simple_sheet(limit_sheet, limit_headers, limit_rows, [13, 14, 14, 13, 14, 17, 14, 14, 11])

    history_sheet = workbook.create_sheet("导入历史")
    history_rows = [[item["trade_date"], item["row_count"], item["original_name"], item["sha256"], item["imported_at"]]
                    for item in store.import_records()]
    _write_simple_sheet(history_sheet, ["交易日期", "股票数", "原文件名", "SHA-256", "导入时间"], history_rows,
                        [13, 11, 42, 68, 22])

    latest_codes = {row["code"] for row in latest_rows}
    market_matches = len(latest_codes & set(market))
    limit_matches = len(latest_codes & set(limits))
    volume_matches = sum(1 for code in latest_codes if volume_history and volume_history.get(code))
    checks = [
        ["报告日源数据行数", len(latest_rows), ">=5000", "通过" if len(latest_rows) >= 5000 else "失败"],
        ["报告输出行数", len(rows), f"应等于 {len(latest_rows)}", "通过" if len(rows) == len(latest_rows) else "失败"],
        ["历史交易日数量", len(dates), "最多读取11日", "提示"],
        ["3日双窗口计算", len(dates), ">=4个交易日", "通过" if len(dates) >= 4 else "失败"],
        ["5日双窗口计算", len(dates), ">=6个交易日", "通过" if len(dates) >= 6 else "失败"],
        ["7日双窗口计算", len(dates), ">=8个交易日", "通过" if len(dates) >= 8 else "失败"],
        ["10日双窗口计算", len(dates), ">=11个交易日", "通过" if len(dates) >= 11 else "失败"],
        ["精确市值匹配", market_matches, f"报告日{len(latest_codes)}只", "通过" if market_matches >= 4000 else "提示（可选Token未设置或无权限）"],
        ["真实成交量匹配", volume_matches, f"报告日{len(latest_codes)}只", "通过" if volume_matches >= 4000 else "提示（未设置Token或成交量历史不足）"],
        ["涨停池总数", len(limits), ">0", "通过" if limits else "失败"],
        ["涨停与源股票匹配", limit_matches, f"涨停池{len(limits)}只", "通过" if limit_matches == len(limits) else "提示"],
    ]
    check_sheet = workbook.create_sheet("数据检查")
    _write_simple_sheet(check_sheet, ["检查项", "实际值", "标准", "结果"], checks, [28, 18, 24, 12])
    for row_index in range(2, len(checks) + 2):
        result = check_sheet.cell(row_index, 4)
        if result.value == "通过":
            result.fill = PatternFill("solid", fgColor=PALE_GREEN)
            result.font = Font(color=GREEN, bold=True)
        elif result.value == "失败":
            result.fill = PatternFill("solid", fgColor=PALE_RED)
            result.font = Font(color=RED, bold=True)

    guide = workbook.create_sheet("使用说明")
    guide.sheet_view.showGridLines = False
    notes = [
        ("报表日期", report_date.isoformat()),
        ("资金单位", "资金净额与差值均为万元；总市值为亿元；封单资金为万元。"),
        ("昨日对比", "按已导入历史中最近两个交易日比较，同一股票缺值时保持空白，不伪造数据。"),
        ("资金流入比率", "资金流入比率=同一资金成分净额/同一窗口成交额×100%；主力、特大单、大单和中单分别计算。"),
        ("比率简表", "首页显示当日和3/5/7/10日主力资金流入比率；分差=含今日比率-不含今日同长度比率，单位为百分点。"),
        ("周期资金观察", "按10日→7日→5日→3日展示主力净额、区间成交额及主力资金流入比率；相邻变化=后一个较短周期比率-前一个较长周期比率；逐步增减值=3日比率-10日比率，并标识越接近今日是否逐步增长。"),
        ("资金成分", "日、3日、5日、7日、10日分别拆分主力、特大单、大单和中单，显示两个窗口净额、增减量、资金流入比率和比率分差。"),
        ("周期递进", "含今日和不含今日分别提供周期递进表，展示10/7/5/3日净额及资金流入比率，以及10→7、7→5、5→3日的增加量和比率变化。"),
        ("市值占比", "净额占总市值=对应窗口主力净额/报告日总市值×100%；成交额占总市值=对应窗口成交额/报告日总市值×100%。"),
        ("10日成交量倍量", "成交量来自Tushare日线真实成交量；每一日倍量=该日成交量/前一交易日成交量。1.5倍至不足2倍为蓝色，2倍至不足3倍为紫色，3倍及以上为黑色。"),
        ("滚动周期", "含今日为最近N个交易日；不含今日为向前平移1个交易日的N日窗口。"),
        ("颜色", "数值增加为红色、减少为绿色，颜色深浅表示变化幅度等级。"),
        ("筛选与排序", "每张分析表均按完整数据区域筛选。点击任一指标排序时，序号、代码、名称和其他指标随整行同步移动；原序号与股票绑定，不重新编号。基础信息列冻结不动。"),
        ("跨表统一排序", f"生成时统一按“{sort_option}”{sort_direction}排列；所有分析页、当日源数据和涨停明细复用同一股票顺序。Excel内手动排序只影响当前页，需要更换跨表顺序时请在软件中重新选择并生成。"),
        ("市值来源", "Tushare daily_basic 的报告日总市值与收盘价。停牌等无当日记录股票保持空白。"),
        ("涨停来源", "东方财富涨停池；包含涨停幅度、几天几板、连板数、封单资金和封板时间。"),
        ("完整性原则", "Tushare为可选补充，未设置时市值留空；涨停数据或源表校验失败时停止生成，不用估算值冒充正式结果。"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row_index, (title, text) in enumerate(notes, start=1):
        guide.cell(row_index, 1, title).font = Font(bold=True, color=NAVY)
        guide.cell(row_index, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row_index].height = 34
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110

    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    progress("正在执行导出文件结构验收...")
    verify = load_workbook(temporary, read_only=True, data_only=False)
    if verify.sheetnames[:4] != ["周期资金观察", "10日成交量倍量", "资金比率简表", "日资金成分对比"]:
        verify.close()
        temporary.unlink(missing_ok=True)
        raise RuntimeError("导出文件工作表结构验收失败。")
    if verify["周期资金观察"].max_row != len(rows) + 2:
        verify.close()
        temporary.unlink(missing_ok=True)
        raise RuntimeError("导出文件行数验收失败。")
    verify.close()
    os.replace(temporary, output)
    progress(f"报表已通过验收：{output}")
    return output
