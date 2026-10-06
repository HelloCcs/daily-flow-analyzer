from __future__ import annotations

import hashlib
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import FlowRow, ParsedWorkbook


EXPECTED_HEADERS = {
    1: "代码",
    2: "名称",
    3: "区间涨跌幅(%)",
    4: "成交额(万元)",
    5: "净流入额(万元)",
    6: "资金流入比率(%)",
    11: "净买入额(万元)",
    17: "净买入额(万元)",
    23: "净买入额(万元)",
    29: "净买入额(万元)",
    33: "起始日期",
    34: "截止日期",
}


class SourceValidationError(ValueError):
    pass


def _number(value: Any) -> float | None:
    if value is None or value == "" or value == "-":
        return None
    if isinstance(value, str):
        value = value.replace(",", "").replace("%", "").strip()
        if not value or value == "-":
            return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_workbook(file_path: str | Path) -> ParsedWorkbook:
    path = Path(file_path).resolve()
    if not path.is_file():
        raise SourceValidationError(f"文件不存在：{path}")
    if path.suffix.lower() != ".xlsx":
        raise SourceValidationError("只支持网站导出的 .xlsx 文件。")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise SourceValidationError(f"Excel 文件无法读取：{exc}") from exc
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        next(rows)
        header = next(rows)
    except StopIteration as exc:
        raise SourceValidationError("工作表为空或缺少两行表头。") from exc

    mismatches = []
    for index, expected in EXPECTED_HEADERS.items():
        actual = str(header[index] or "").strip()
        if actual != expected:
            mismatches.append(f"第{index + 1}列应为“{expected}”，实际为“{actual}”")
    if mismatches:
        raise SourceValidationError("不是受支持的网站资金流向表：" + "；".join(mismatches[:3]))

    parsed: list[FlowRow] = []
    dates: set[date] = set()
    codes: set[str] = set()
    invalid_dates = 0
    invalid_ranges: set[tuple[date | None, date | None]] = set()
    missing_numeric = 0
    relation_errors = 0
    for excel_row, row in enumerate(rows, start=3):
        code = str(row[1] or "").strip().upper()
        if not code:
            continue
        if code in codes:
            raise SourceValidationError(f"股票代码重复：{code}（约第 {excel_row} 行）")
        codes.add(code)
        start_date, end_date = _date(row[33]), _date(row[34])
        if not start_date or not end_date or start_date != end_date:
            invalid_dates += 1
            invalid_ranges.add((start_date, end_date))
            continue
        dates.add(end_date)
        values = [_number(row[i]) for i in (3, 4, 5, 6, 11, 17, 23, 29)]
        if values[2] is None or values[4] is None or values[5] is None or values[6] is None:
            missing_numeric += 1
        main, super_net, big_net = values[2], values[4], values[5]
        if main is not None and super_net is not None and big_net is not None:
            if abs(main - super_net - big_net) > 0.05:
                relation_errors += 1
        parsed.append(
            FlowRow(
                trade_date=end_date,
                code=code,
                name=str(row[2] or "").strip(),
                pct_change=values[0],
                turnover=values[1],
                main_net=values[2],
                main_ratio=values[3],
                super_net=values[4],
                big_net=values[5],
                mid_net=values[6],
                small_net=values[7],
            )
        )
    workbook.close()

    if invalid_dates:
        ranges = "、".join(
            f"{start or '无效日期'} 至 {end or '无效日期'}"
            for start, end in sorted(invalid_ranges, key=lambda item: (str(item[0]), str(item[1])))
        )
        raise SourceValidationError(
            f"该文件是累计区间表（检测到 {ranges}），不是单日表。"
            f"请把网站的起始日期和截止日期都设为同一个交易日后重新导出；"
            f"本次有 {invalid_dates} 行被拒绝，未写入历史库。"
        )
    if len(dates) != 1:
        raise SourceValidationError(f"一个文件必须只包含一个交易日，检测到 {len(dates)} 个日期。")
    if len(parsed) < 5000:
        raise SourceValidationError(f"有效股票仅 {len(parsed)} 行，少于完整样表的 5000 行，已拒绝导入。")
    if missing_numeric > max(30, len(parsed) // 100):
        raise SourceValidationError(f"关键资金字段缺失 {missing_numeric} 行，超过允许范围。")
    if relation_errors:
        raise SourceValidationError(f"有 {relation_errors} 行不满足主力净额=特大单净额+大单净额。")

    warnings = []
    if missing_numeric:
        warnings.append(f"{missing_numeric} 只股票当天无交易或关键金额为空，报表中保留为空。")
    return ParsedWorkbook(path, next(iter(dates)), _hash(path), tuple(parsed), tuple(warnings))
